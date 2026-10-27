#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
create_bidv_minutes.py
Công cụ tự động hóa toàn trình (Pipeline) tạo Biên bản cuộc họp / Thông báo kết luận chuẩn BIDV:
1. Đọc nội dung file transcript / ghi chép (.docx, .txt, .md).
2. Gọi trực tiếp mô hình cao cấp 'gemini-argon' qua gemini-stealth-bridge (http://127.0.0.1:20133/v1).
3. Đóng gói kết quả thành file Word (.docx) và PDF theo đúng 100% quy chuẩn skill bidv-meeting-minutes:
   - Khổ giấy: US Letter (8.5 x 11.0 in)
   - Căn lề: Top 0.75", Bottom 0.75", Left 1.0", Right 0.75"
   - Phông chữ: Times New Roman 13pt
   - Header bảng không viền (Quốc hiệu Tiêu ngữ & Cơ quan ban hành)
   - Cấu trúc: Dẫn nhập -> 1. Bối cảnh (có (i)... (v)...) -> 2. Chỉ đạo (2.1 chung, 2.2 cụ thể 3 gạch đầu dòng) -> Đoạn kết -> Footer chữ ký không tách trang.
   - Lọc bỏ sạch sẽ 100% các ký tự markdown (*, **, ---, ===) và thẻ HTML (<table...>, <tr>, <td>, <p>, <br>...).
   - Chuyển đổi toàn bộ in đậm/in nghiêng thành thuộc tính Font Run thực sự trong Word.
"""

import argparse
import json
import os
import re
import subprocess
import sys
import time
import urllib.request
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

DEFAULT_BRIDGE_URL = "http://127.0.0.1:20133/v1/chat/completions"
DEFAULT_MODEL = "gemini-argon"
API_KEY = "public"

def set_cell_borderless(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        r'<w:tcBorders {}>'
        r'<w:top w:val="none"/><w:left w:val="none"/>'
        r'<w:bottom w:val="none"/><w:right w:val="none"/>'
        r'</w:tcBorders>'.format(nsdecls('w'))
    )
    tcPr.append(tcBorders)

def set_run_font(run, size_pt=13.0, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        r'<w:rFonts {} w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'.format(
            nsdecls("w")
        )
    )
    rPr.append(rFonts)

def parse_markdown_runs(text):
    """
    Tách chuỗi chứa markdown (**đậm**, *nghiêng*) thành danh sách (run_text, is_bold, is_italic)
    Loại bỏ hoàn toàn các ký tự '*', '**' để text trong Word sạch 100%.
    """
    # Regex tìm token **...**, *...* hoặc text bình thường
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|[^*]+)")
    tokens = pattern.findall(text)
    runs = []
    for t in tokens:
        if t.startswith("**") and t.endswith("**") and len(t) >= 4:
            clean = t[2:-2]
            runs.append((clean, True, False))
        elif t.startswith("*") and t.endswith("*") and len(t) >= 2:
            clean = t[1:-1]
            runs.append((clean, False, True))
        else:
            runs.append((t, False, False))
    return runs

def add_runs_to_paragraph(paragraph, text, default_size=13.0, force_bold=False, force_italic=False):
    """Thêm các Run văn bản vào Paragraph đã bóc tách định dạng đậm/nghiêng sạch sẽ."""
    # Xóa sạch mọi thẻ HTML còn sót
    clean_text = text.replace("&nbsp;", " ").replace("<br/>", " ").replace("<br>", " ")
    clean_text = re.sub(r"<[^>]+>", "", clean_text)
    
    parsed = parse_markdown_runs(clean_text)
    for run_text, b, it in parsed:
        if not run_text:
            continue
        run = paragraph.add_run(run_text)
        set_run_font(
            run, 
            size_pt=default_size, 
            bold=(True if force_bold else b), 
            italic=(True if force_italic else it)
        )

def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        doc = docx.Document(file_path)
        lines = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                lines.append(t)
        for tbl in doc.tables:
            for row in tbl.rows:
                r_str = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip())
                if r_str:
                    lines.append(r_str)
        return "\n".join(lines)
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

def call_gemini_argon(transcript_text, model=DEFAULT_MODEL, bridge_url=DEFAULT_BRIDGE_URL):
    system_prompt = """Bạn là Thư ký Cấp cao của Ngân hàng TMCP Đầu tư và Phát triển Việt Nam (BIDV).
Nhiệm vụ của bạn là đọc kỹ bản ghi chép chi tiết, phân vai của cuộc họp/hội nghị và soạn thảo thành nội dung **BIÊN BẢN CUỘC HỌP** (hoặc THÔNG BÁO KẾT LUẬN) chính thức theo đúng chuẩn mực thể thức văn bản hành chính của BIDV (theo skill bidv-meeting-minutes).

QUY ĐỊNH ĐẶC BIỆT VỀ ĐỊNH DẠNG:
- TUYỆT ĐỐI KHÔNG xuất các thẻ HTML như <table>, <tr>, <td>, <p>, <br>, <b>, <i>, <div>...
- TUYỆT ĐỐI KHÔNG tạo bảng Quốc hiệu, Tiêu ngữ, số hiệu hay Bảng Header đầu trang (hệ thống Word sẽ tự động chèn).
- TUYỆT ĐỐI KHÔNG tạo bảng Nơi nhận hay Chữ ký cuối trang (hệ thống Word sẽ tự động chèn).
- TUYỆT ĐỐI KHÔNG dùng các đường kẻ phân cách thô '---' hoặc '==='.
- Chỉ sử dụng định dạng văn bản chuẩn: in đậm tiêu đề bằng **Tiêu đề**, không dùng ký tự trang trí lạ.

Yêu cầu cấu trúc nội dung bắt buộc:
1. PHẦN DẪN NHẬP (2 đoạn văn mở đầu chuẩn BIDV):
   - Đoạn 1: Thời gian, địa điểm, người chủ trì (TGĐ Hoàng Việt Hùng và Phó Chủ tịch Tencent Danny Hong), thành phần tham dự chi tiết (5 Phó TGĐ BIDV gồm bà Nguyễn Thị Quỳnh Giao, ông Phan Thanh Hải, ông Lại Tiến Quân, ông Vũ Hoàng Dương, ông Lê Huy Hoàng, Phó Trưởng khối Nguyễn Hữu Tiến, Phó GĐVP Ngô Minh Hoàng và 20 Lãnh đạo Ban/TT/CN; Lãnh đạo cấp cao Tencent).
   - Đoạn 2: "Sau khi nghe đại diện Tencent trình bày tổng quan... ý kiến thảo luận của các đơn vị... Tổng Giám đốc BIDV Hoàng Việt Hùng kết luận và thống nhất phân công triển khai như sau:"
2. MỤC 1. BỐI CẢNH CUỘC HỌP:
   - Trình bày trong đúng 3 đoạn văn hoàn chỉnh.
   - Đoạn 1: Bối cảnh, mục đích triệu tập cuộc họp.
   - Đoạn 2: Thực trạng giải pháp, diễn biến trao đổi của hai bên.
   - Đoạn 3: Kết quả rà soát và các bài toán/điểm nghẽn trọng tâm, đánh số thứ tự nối tiếp dạng: (i) ...; (ii) ...; (iii) ...; (iv) ...; (v) ...
3. MỤC 2. NỘI DUNG THỐNG NHẤT VÀ CHỈ ĐẠO TRIỂN KHAI:
   - Mục 2.1. Định hướng và chỉ đạo chung: 4 đến 5 luận điểm nguyên tắc xuyên suốt, mỗi luận điểm mở đầu bằng cụm động từ mạnh in đậm có dấu hai chấm:
     **1. Chuyển đổi công nghệ gắn liền với hiệu quả thực chất: **...
     **2. Đặt an toàn hệ thống, bảo mật dữ liệu và tuân thủ pháp luật lên hàng đầu: **...
     **3. Tăng tốc hoàn thiện kiến trúc AI Platform và khung quản trị AI (AI Governance): **...
     **4. Khai thác tối đa lợi thế hệ sinh thái mở và mạng lưới thanh toán xuyên biên giới: **...
     **5. Tăng cường phối hợp chặt chẽ và gắn kết trách nhiệm đầu mối: **...
   - Mục 2.2. Nội dung triển khai cụ thể: nội dung - đơn vị thực hiện - thời hạn:
     Chia thành các đầu mục công việc in đậm "1. Về [Chủ đề]:", "2. Về [Chủ đề]:"...
     Dưới mỗi đầu mục BẮT BUỘC có đúng 3 gạch đầu dòng:
     - **- Nội dung chỉ đạo và thống nhất: ** Nêu chi tiết nội dung nghiệp vụ, kỹ thuật, giải pháp.
     - **- Đơn vị thực hiện: ** Ghi rõ Giao [Tên Ban/Trung tâm] (chủ trì / làm đầu mối), phối hợp [Các Ban/TT]...
     - **- Thời hạn hoàn thành: ** Ghi rõ mốc thời gian hoàn thành (đặt trong dấu *in nghiêng*).
4. ĐOẠN KẾT VĂN BẢN:
   - "Trên đây là nội dung biên bản kết luận và thống nhất phân công triển khai tại Phiên họp cấp cao song phương ngày 27/10/2026 giữa Ban Lãnh đạo BIDV và Ban Lãnh đạo Tập đoàn Tencent... Văn phòng BIDV trân trọng thông báo để các đơn vị biết và nghiêm túc triển khai thực hiện./." """

    user_prompt = f"""Dưới đây là toàn bộ bản ghi chép chi tiết cuộc họp:
---
{transcript_text}
---

Hãy sử dụng năng lực suy luận sâu của mô hình `{model}` để phân tích thấu đáo, chuẩn hóa dữ liệu nhân sự, chức danh và biên soạn Biên bản cuộc họp chất lượng cao nhất theo đúng các yêu cầu trên."""

    req_data = {
        "model": model,
        "stream": True,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]
    }

    print(f"[*] Đang gửi dữ liệu ({len(transcript_text)} ký tự) tới mô hình [{model}] qua {bridge_url}...")
    req = urllib.request.Request(
        bridge_url,
        data=json.dumps(req_data).encode("utf-8"),
        headers={
            "Authorization": f"Bearer {API_KEY}",
            "Content-Type": "application/json"
        }
    )

    start = time.time()
    chunks = []
    total_chars = 0

    with urllib.request.urlopen(req, timeout=300) as resp:
        for line in resp:
            line = line.decode("utf-8").strip()
            if not line or line.startswith(": ping"):
                continue
            if line.startswith("data: "):
                payload = line[6:]
                if payload == "[DONE]":
                    break
                try:
                    data = json.loads(payload)
                    delta = data["choices"][0]["delta"]
                    content = delta.get("content", "")
                    if content:
                        chunks.append(content)
                        total_chars += len(content)
                        sys.stdout.write(f"\r    -> Mô hình [{model}] đang sinh nội dung: {total_chars} ký tự ({time.time() - start:.1f}s)...")
                        sys.stdout.flush()
                except Exception:
                    pass

    print(f"\n[+] Hoàn thành sinh nội dung từ [{model}] trong {time.time() - start:.1f}s! (Tổng: {total_chars} ký tự)")
    return "".join(chunks)

def is_ignorable_line(line):
    """Kiểm tra và loại bỏ các dòng rác markdown, HTML hoặc dòng chào hỏi/tiêu đề lặp."""
    l = line.strip()
    if not l:
        return True
    # Bỏ các thẻ HTML
    if l.startswith("<table") or l.startswith("</table") or l.startswith("<tr") or l.startswith("</tr") or \
       l.startswith("<td") or l.startswith("</td") or l.startswith("<p") or l.startswith("</p") or \
       l.startswith("<div") or l.startswith("</div"):
        return True
    # Bỏ các đường kẻ markdown
    if re.match(r"^[-=_*]{3,}$", l):
        return True
    # Bỏ lời dẫn nhập của AI
    if l.startswith("Dưới đây là") or l.startswith("Sau đây là") or l.startswith("Toàn văn Biên bản"):
        return True
    # Bỏ các tiêu đề lặp lại vì đã có trong Header Table và Title Block
    if any(h in l for h in [
        "BIÊN BẢN CUỘC HỌP", "BIÊN BẢN HỌP CẤP CAO", "THÔNG BÁO KẾT LUẬN",
        "CỘNG HÒA XÃ HỘI", "CỘNG HOÀ XÃ HỘI", "Độc lập - Tự do", 
        "NGÂN HÀNG TMCP ĐẦU TƯ", "TẬP ĐOÀN TENCENT (TENCENT",
        "Số: 2710", "Số:          /BB", "Số:          /TB",
        "V/v: Thúc đẩy hợp tác chiến lược", "V/v Tọa đàm cấp cao"
    ]):
        return True
    return False

def set_table_styling(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders {}>'
        r'<w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>'
        r'<w:left w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>'.format(nsdecls('w'))
    )
    tblPr.append(borders)

def build_bidv_docx(markdown_content, output_path):
    print(f"[*] Đang đóng gói file Word theo chuẩn thể thức BIDV (Lọc sạch 100% Markdown/HTML): {output_path}...")
    doc = docx.Document()

    # 1. Page margins
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.75)

    # 2. Header Table (2 cột không viền)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_left = tbl.cell(0, 0)
    c_right = tbl.cell(0, 1)
    c_left.width = Inches(3.25)
    c_right.width = Inches(3.50)
    set_cell_borderless(c_left)
    set_cell_borderless(c_right)

    # Left: Đơn vị ban hành
    p = c_left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("NGÂN HÀNG TMCP ĐẦU TƯ\nVÀ PHÁT TRIỂN VIỆT NAM")
    set_run_font(r, size_pt=12.0, bold=True)

    p_l2 = c_left.add_paragraph()
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l2.paragraph_format.space_after = Pt(2)
    r = p_l2.add_run("————————")
    set_run_font(r, size_pt=11.0)

    p_so = c_left.add_paragraph()
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_so.paragraph_format.space_after = Pt(0)
    r = p_so.add_run("Số:   2710   /BB-BIDV")
    set_run_font(r, size_pt=12.0)

    # Right: Quốc hiệu Tiêu ngữ & Ngày tháng
    p_r1 = c_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.line_spacing = 1.15
    p_r1.paragraph_format.space_after = Pt(2)
    r = p_r1.add_run("CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    set_run_font(r, size_pt=12.0, bold=True)

    p_r2 = c_right.add_paragraph()
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r2.paragraph_format.space_after = Pt(2)
    r = p_r2.add_run("Độc lập - Tự do - Hạnh phúc")
    set_run_font(r, size_pt=12.5, bold=True)

    p_r3 = c_right.add_paragraph()
    p_r3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r3.paragraph_format.space_after = Pt(4)
    r = p_r3.add_run("————————————")
    set_run_font(r, size_pt=11.0)

    p_date = c_right.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_after = Pt(0)
    r = p_date.add_run("Thâm Quyến, ngày 27 tháng 10 năm 2026")
    set_run_font(r, size_pt=12.0, italic=True)

    # Spacing after header
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(0)

    # 3. Title Block
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(16)
    p_t.paragraph_format.space_after = Pt(6)
    r = p_t.add_run("BIÊN BẢN CUỘC HỌP")
    set_run_font(r, size_pt=15.0, bold=True)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(4)
    r = p_sub.add_run("Phiên họp Cấp cao Song phương giữa Ban Lãnh đạo BIDV và Ban Lãnh đạo Tập đoàn Tencent\nvề Hợp tác Chiến lược Toàn diện trong lĩnh vực Chuyển đổi số, AI, Cloud và Thanh toán")
    set_run_font(r, size_pt=13.5, bold=True)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(14)
    r = p_line.add_run("————————————")
    set_run_font(r, size_pt=12.0)

    # 4. Group lines into blocks (paragraphs or tables)
    raw_lines = markdown_content.splitlines()
    i = 0
    num_lines = len(raw_lines)

    while i < num_lines:
        line = raw_lines[i].strip()

        if is_ignorable_line(line):
            i += 1
            continue

        if "Nơi nhận:" in line or "ĐẠI DIỆN ĐỒNG CHỦ TRÌ" in line or "CHỮ KÝ" in line:
            break

        # Check if line is part of a Markdown Table
        if line.startswith("|") and line.endswith("|"):
            table_lines = []
            while i < num_lines and raw_lines[i].strip().startswith("|") and raw_lines[i].strip().endswith("|"):
                tl = raw_lines[i].strip()
                if not re.match(r"^\|\s*[:\-]+\s*\|", tl):
                    cells = [c.strip() for c in tl.strip("|").split("|")]
                    table_lines.append(cells)
                i += 1

            if table_lines:
                num_cols = max(len(row) for row in table_lines)
                wtbl = doc.add_table(rows=len(table_lines), cols=num_cols)
                wtbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_table_styling(wtbl)

                for r_idx, row_cells in enumerate(table_lines):
                    is_header_row = (r_idx == 0)
                    for c_idx in range(num_cols):
                        cell_val = row_cells[c_idx] if c_idx < len(row_cells) else ""
                        wcell = wtbl.cell(r_idx, c_idx)
                        cp = wcell.paragraphs[0]
                        cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(2)
                        cp.paragraph_format.line_spacing = 1.15
                        add_runs_to_paragraph(cp, cell_val, default_size=11.0 if not is_header_row else 11.5, force_bold=is_header_row)

                # Add small space after table
                p_after_tbl = doc.add_paragraph()
                p_after_tbl.paragraph_format.space_before = Pt(6)
                p_after_tbl.paragraph_format.space_after = Pt(0)
            continue

        # Tiêu đề mục lớn: 1. BỐI CẢNH CUỘC HỌP, 2. NỘI DUNG THỐNG NHẤT VÀ CHỈ ĐẠO TRIỂN KHAI
        if line.startswith("## ") or (len(line) > 2 and line[0] in "12345" and line[1] == "." and " " in line[:4] and not line.startswith("2.1") and not line.startswith("2.2")):
            clean = line.lstrip("#").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_runs_to_paragraph(p, clean, default_size=13.0, force_bold=True)

        # Tiêu đề mục con: 2.1. Định hướng..., 2.2. Nội dung triển khai...
        elif line.startswith("### ") or line.startswith("2.1") or line.startswith("2.2"):
            clean = line.lstrip("#").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_runs_to_paragraph(p, clean, default_size=13.0, force_bold=True)

        # 3 gạch đầu dòng chuẩn của BIDV:
        elif any(line.startswith(prefix) for prefix in [
            "- **- Nội dung", "- **- Đơn vị", "- **- Thời hạn",
            "- Nội dung chỉ đạo", "- Đơn vị thực hiện", "- Thời hạn hoàn thành",
            "**- Nội dung", "**- Đơn vị", "**- Thời hạn",
            "* **- Nội dung", "* **- Đơn vị", "* **- Thời hạn",
            "* - Nội dung", "* - Đơn vị", "* - Thời hạn"
        ]):
            clean = re.sub(r"^[* \-]+", "", line).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(0.0)
            p.paragraph_format.space_after = Pt(4)

            if ":" in clean:
                label_part, val_part = clean.split(":", 1)
                label_clean = "- " + label_part.replace("**", "").replace("*", "").strip() + ": "
                
                r_lbl = p.add_run(label_clean)
                set_run_font(r_lbl, size_pt=13.0, bold=True)
                
                is_time = "Thời hạn" in label_clean
                if is_time:
                    val_clean = val_part.replace("**", "").replace("*", "").strip()
                    r_val = p.add_run(val_clean)
                    set_run_font(r_val, size_pt=13.0, italic=True)
                else:
                    add_runs_to_paragraph(p, val_part, default_size=13.0)
            else:
                add_runs_to_paragraph(p, "- " + clean, default_size=13.0)

        # Đầu mục công việc trong 2.2: 1. Về ...:, 2. Về ...:
        elif re.match(r"^(\*\*)?[0-9]+\.\s+Về", line):
            clean = line.replace("**", "").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.keep_with_next = True
            add_runs_to_paragraph(p, clean, default_size=13.0, force_bold=True)

        # Các luận điểm nguyên tắc trong 2.1: 1. Động từ mạnh: ..., 2. Động từ mạnh: ...
        elif re.match(r"^[0-9]+\.\s+", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)

            if ":" in line:
                pre, post = line.split(":", 1)
                r_pre = p.add_run(pre.replace("**", "").replace("*", "").strip() + ": ")
                set_run_font(r_pre, size_pt=13.0, bold=True)
                add_runs_to_paragraph(p, post, default_size=13.0)
            else:
                add_runs_to_paragraph(p, line, default_size=13.0)

        # Các dòng gạch đầu dòng thông thường
        elif line.startswith("* ") or line.startswith("- "):
            clean = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(0.0)
            p.paragraph_format.space_after = Pt(3)
            
            r_bullet = p.add_run("- ")
            set_run_font(r_bullet, size_pt=13.0)
            add_runs_to_paragraph(p, clean, default_size=13.0)

        # Đoạn văn xuôi thông thường
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)
            add_runs_to_paragraph(p, line, default_size=13.0)

        i += 1

    # 5. Footer Table: Bảng Nơi nhận và Chữ ký cuối trang (1 dòng, 2 cột không viền)
    f_tbl = doc.add_table(rows=1, cols=2)
    f_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_tbl.autofit = False

    trPr = f_tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))

    f_left = f_tbl.cell(0, 0)
    f_right = f_tbl.cell(0, 1)
    f_left.width = Inches(3.50)
    f_right.width = Inches(3.25)
    set_cell_borderless(f_left)
    set_cell_borderless(f_right)

    # Nơi nhận
    p_nn = f_left.paragraphs[0]
    p_nn.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_nn.paragraph_format.space_after = Pt(2)
    r = p_nn.add_run("Nơi nhận:")
    set_run_font(r, size_pt=12.0, bold=True, italic=True)

    items = [
        "- Lãnh đạo tham dự cuộc họp;",
        "- Ban Lãnh đạo Tập đoàn Tencent;",
        "- Các Khối, Ban, Trung tâm thuộc BIDV;",
        "- Lưu: VT, VP, Tổ Thư ký đoàn."
    ]
    for it in items:
        p_it = f_left.add_paragraph()
        p_it.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_it.paragraph_format.space_after = Pt(1)
        r = p_it.add_run(it)
        set_run_font(r, size_pt=11.0)

    # Chữ ký
    p_sg = f_right.paragraphs[0]
    p_sg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sg.paragraph_format.line_spacing = 1.15
    p_sg.paragraph_format.space_after = Pt(2)
    r = p_sg.add_run("ĐẠI DIỆN ĐỒNG CHỦ TRÌ PHIÊN HỌP\n")
    set_run_font(r, size_pt=12.5, bold=True)
    r = p_sg.add_run("TỔNG GIÁM ĐỐC BIDV")
    set_run_font(r, size_pt=12.0, bold=True)

    for _ in range(5):
        p_b = f_right.add_paragraph()
        p_b.paragraph_format.space_after = Pt(0)
        p_b.paragraph_format.line_spacing = 1.0

    p_nm = f_right.add_paragraph()
    p_nm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nm.paragraph_format.space_after = Pt(0)
    r = p_nm.add_run("HOÀNG VIỆT HÙNG")
    set_run_font(r, size_pt=12.0, bold=True)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"[+] Xuất thành công file DOCX sạch: {output_path}")

    # Convert to PDF
    pdf_path = os.path.splitext(output_path)[0] + ".pdf"
    try:
        subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", output_path, "--outdir", os.path.dirname(output_path)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if os.path.exists(pdf_path):
            print(f"[+] Xuất thành công file PDF sạch: {pdf_path}")
    except Exception:
        pass

def main():
    parser = argparse.ArgumentParser(description="Tạo biên bản họp chuẩn BIDV sử dụng mô hình gemini-argon")
    parser.add_argument("input_file", help="Đường dẫn file transcript đầu vào (.docx, .txt, .md)")
    parser.add_argument("-o", "--output", help="Đường dẫn file .docx đầu ra", default=None)
    parser.add_argument("--model", help="Tên model Gemini muốn chỉ định (mặc định: gemini-argon)", default=DEFAULT_MODEL)
    args = parser.parse_args()

    input_path = os.path.abspath(args.input_file)
    if not os.path.exists(input_path):
        print(f"[!] Lỗi: Không tìm thấy file {input_path}")
        sys.exit(1)

    if not args.output:
        base_name = os.path.splitext(os.path.basename(input_path))[0]
        output_path = os.path.join(os.path.expanduser("~/Downloads"), f"{base_name}_Bien_ban_chuan_BIDV_argon.docx")
    else:
        output_path = os.path.abspath(args.output)

    print("=" * 70)
    print(" PIPELINE TẠO BIÊN BẢN HỌP BIDV VỚI MÔ HÌNH CAO CẤP: " + args.model.upper())
    print("=" * 70)
    print(f"[1/3] Đang đọc dữ liệu từ: {input_path}")
    transcript = extract_text_from_file(input_path)
    print(f"      -> Đã nạp {len(transcript)} ký tự.")

    print(f"[2/3] Gọi AI Gemini [{args.model}] suy luận sâu & soạn thảo...")
    minutes_md = call_gemini_argon(transcript, model=args.model)

    print(f"[3/3] Đóng gói văn bản theo thể thức hành chính chuẩn BIDV...")
    build_bidv_docx(minutes_md, output_path)

    print("=" * 70)
    print(f"[DONE] Hoàn tất toàn bộ quy trình! File đầu ra:")
    print(f" -> Word: {output_path}")
    pdf_path = os.path.splitext(output_path)[0] + ".pdf"
    if os.path.exists(pdf_path):
        print(f" -> PDF : {pdf_path}")
    print("=" * 70)

if __name__ == "__main__":
    main()
