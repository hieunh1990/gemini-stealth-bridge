#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
bidv_engine.py
Engine sinh Biên bản cuộc họp & Thông báo kết luận chuẩn thể thức BIDV
Tuân thủ 100% skill bidv-meeting-minutes:
- Căn lề: Top 0.75", Bottom 0.75", Left 1.0", Right 0.75" (US Letter)
- Font: Times New Roman 13pt
- Header & Footer tables borderless (Quốc hiệu, Tên ngân hàng, Nơi nhận, Chữ ký cantSplit)
- Cấu trúc: 2 đoạn dẫn nhập -> 1. Bối cảnh (có (i)...(v)...) -> 2. Chỉ đạo (2.1 chung, 2.2 cụ thể 3 gạch đầu dòng) -> Đoạn kết
- Tuyệt đối không chèn bảng ma trận trong thân văn bản
- Bóc tách toàn bộ Markdown sang native Word Run (Bold, Italic)
"""

import json
import os
import re
import subprocess
import time
import requests
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Inches, Pt, RGBColor

GSB_URL = os.environ.get("GSB_URL", "http://gsb:20133/v1/chat/completions")
API_KEY = os.environ.get("GSB_API_KEY", "public")
DEFAULT_MODEL = os.environ.get("DEFAULT_MODEL", "gemini-argon")

def set_cell_borderless(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        r'<w:tcBorders {}>'
        r'<w:top w:val="none"/><w:left w:val="none"/>'
        r'<w:bottom w:val="none"/><w:right w:val="none"/>'
        r'</w:tcBorders>'.format(nsdecls('w'))
    )
    tcPr.append(tcBorders)

def set_run_font(run, size_pt=13.0, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(
        r'<w:rFonts {} w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>'.format(
            nsdecls("w")
        )
    )
    rPr.append(rFonts)

def parse_markdown_runs(text):
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|[^*]+)")
    tokens = pattern.findall(text)
    runs = []
    for t in tokens:
        if t.startswith("**") and t.endswith("**") and len(t) >= 4:
            runs.append((t[2:-2], True, False))
        elif t.startswith("*") and t.endswith("*") and len(t) >= 2:
            runs.append((t[1:-1], False, True))
        else:
            runs.append((t, False, False))
    return runs

def add_runs_to_paragraph(paragraph, text, default_size=13.0, force_bold=False, force_italic=False):
    clean_text = text.replace("&nbsp;", " ").replace("<br/>", " ").replace("<br>", " ")
    clean_text = re.sub(r"<[^>]+>", "", clean_text)
    parsed = parse_markdown_runs(clean_text)
    for run_text, b, it in parsed:
        if not run_text:
            continue
        run = paragraph.add_run(run_text)
        set_run_font(
            run, 
            size_pt=default_size, 
            bold=(True if force_bold else b), 
            italic=(True if force_italic else it)
        )

def extract_text_from_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".docx":
        doc = docx.Document(file_path)
        lines = []
        for p in doc.paragraphs:
            t = p.text.strip()
            if t:
                lines.append(t)
        for tbl in doc.tables:
            for row in tbl.rows:
                r_str = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip())
                if r_str:
                    lines.append(r_str)
        return "\n".join(lines)
    else:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

def call_ai(transcript_text, model=DEFAULT_MODEL, doc_type="BB"):
    type_name = "BIÊN BẢN CUỘC HỌP" if doc_type == "BB" else "THÔNG BÁO KẾT LUẬN"
    action_heading = "2. NỘI DUNG THỐNG NHẤT VÀ CHỈ ĐẠO TRIỂN KHAI" if doc_type == "BB" else "2. CHỈ ĐẠO"
    
    system_prompt = f"""Bạn là Thư ký Cấp cao của Ngân hàng TMCP Đầu tư và Phát triển Việt Nam (BIDV).
Nhiệm vụ của bạn là đọc kỹ bản ghi chép chi tiết, phân vai của cuộc họp/hội nghị và soạn thảo thành nội dung **{type_name}** chính thức theo đúng chuẩn mực thể thức văn bản hành chính của BIDV (theo skill bidv-meeting-minutes).

QUY ĐỊNH ĐẶC BIỆT VỀ ĐỊNH DẠNG:
- TUYỆT ĐỐI KHÔNG xuất các thẻ HTML như <table>, <tr>, <td>, <p>, <br>, <b>, <i>, <div>...
- TUYỆT ĐỐI KHÔNG tạo bảng Quốc hiệu, Tiêu ngữ, số hiệu hay Bảng Header đầu trang (hệ thống Word sẽ tự động chèn).
- TUYỆT ĐỐI KHÔNG tạo bảng Nơi nhận hay Chữ ký cuối trang (hệ thống Word sẽ tự động chèn).
- TUYỆT ĐỐI KHÔNG chèn bảng biểu ma trận nhiệm vụ vào phần thân văn bản.
- TUYỆT ĐỐI KHÔNG dùng các đường kẻ phân cách thô '---' hoặc '==='.

Yêu cầu cấu trúc nội dung bắt buộc:
1. PHẦN DẪN NHẬP (2 đoạn văn mở đầu chuẩn BIDV):
   - Đoạn 1: Thời gian, địa điểm, người chủ trì (Họ tên, chức vụ), thành phần tham dự chi tiết (Lãnh đạo BIDV, các Ban/TT/CN; Lãnh đạo cấp cao đối tác/đơn vị).
   - Đoạn 2: "Sau khi nghe báo cáo/trình bày... ý kiến thảo luận của các đơn vị... [Họ tên người chủ trì] kết luận, chỉ đạo (hoặc thống nhất phân công triển khai) như sau:"
2. MỤC 1. BỐI CẢNH CUỘC HỌP:
   - Trình bày trong đúng 3 đoạn văn hoàn chỉnh.
   - Đoạn 1: Bối cảnh, mục đích triệu tập cuộc họp.
   - Đoạn 2: Thực trạng giải pháp, diễn biến trao đổi của các bên.
   - Đoạn 3: Kết quả rà soát và các bài toán/điểm nghẽn trọng tâm, đánh số thứ tự nối tiếp dạng: (i) ...; (ii) ...; (iii) ...; (iv) ...; (v) ...
3. MỤC {action_heading}:
   - Mục 2.1. Định hướng và chỉ đạo chung: 4 đến 5 luận điểm nguyên tắc xuyên suốt, mỗi luận điểm mở đầu bằng cụm động từ mạnh in đậm có dấu hai chấm:
     **1. [Cụm động từ mạnh]: **...
     **2. [Cụm động từ mạnh]: **...
     **3. [Cụm động từ mạnh]: **...
   - Mục 2.2. Nội dung triển khai cụ thể: nội dung - đơn vị thực hiện - thời hạn:
     Chia thành các đầu mục công việc in đậm "1. Về [Chủ đề]:", "2. Về [Chủ đề]:"...
     Dưới mỗi đầu mục BẮT BUỘC có đúng 3 gạch đầu dòng:
     - **- Nội dung chỉ đạo và thống nhất: ** Nêu chi tiết nội dung nghiệp vụ, kỹ thuật, giải pháp.
     - **- Đơn vị thực hiện: ** Ghi rõ Giao [Tên Ban/Trung tâm] (chủ trì / làm đầu mối), phối hợp [Các Ban/TT]...
     - **- Thời hạn hoàn thành: ** Ghi rõ mốc thời gian hoàn thành (đặt trong dấu *in nghiêng*).
4. ĐOẠN KẾT VĂN BẢN:
   - "Trên đây là ý kiến kết luận chỉ đạo của [Chức vụ + Họ tên] (hoặc nội dung biên bản kết luận và thống nhất phân công triển khai)... Đơn vị ban hành (hoặc Đơn vị đầu mối) trân trọng thông báo để các đơn vị biết và nghiêm túc triển khai thực hiện./." """

    user_prompt = f"""Dưới đây là toàn bộ bản ghi chép chi tiết cuộc họp:
---
{transcript_text}
---

Hãy sử dụng năng lực suy luận sâu của mô hình `{model}` để phân tích thấu đáo, chuẩn hóa dữ liệu nhân sự, chức danh và biên soạn {type_name} chất lượng cao nhất theo đúng các yêu cầu trên."""

    payload = {
        "model": model,
        "stream": True,
        "messages": [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt}
        ]
    }

    headers = {
        "Authorization": f"Bearer {API_KEY}",
        "Content-Type": "application/json"
    }

    resp = requests.post(GSB_URL, json=payload, headers=headers, stream=True, timeout=300)
    resp.raise_for_status()

    chunks = []
    for line in resp.iter_lines():
        if not line:
            continue
        line_str = line.decode("utf-8").strip()
        if line_str.startswith(": ping"):
            continue
        if line_str.startswith("data: "):
            data_str = line_str[6:]
            if data_str == "[DONE]":
                break
            try:
                chunk_json = json.loads(data_str)
                delta = chunk_json["choices"][0]["delta"]
                c = delta.get("content", "")
                if c:
                    chunks.append(c)
            except Exception:
                pass

    return "".join(chunks)

call_gemini = call_ai

def is_ignorable_line(line):
    l = line.strip()
    if not l:
        return True
    if any(l.startswith(p) for p in ["<table", "</table", "<tr", "</tr", "<td", "</td", "<p", "</p", "<div", "</div"]):
        return True
    if re.match(r"^[-=_*]{3,}$", l):
        return True
    if any(l.startswith(p) for p in ["Dưới đây là", "Sau đây là", "Toàn văn Biên bản"]):
        return True
    if any(h in l for h in [
        "BIÊN BẢN CUỘC HỌP", "BIÊN BẢN HỌP CẤP CAO", "THÔNG BÁO KẾT LUẬN",
        "CỘNG HÒA XÃ HỘI", "CỘNG HOÀ XÃ HỘI", "Độc lập - Tự do", 
        "NGÂN HÀNG TMCP ĐẦU TƯ", "Số: ", "Số:   ", "V/v: ", "V/v "
    ]):
        return True
    return False

def build_docx_and_pdf(
    ai_text, 
    docx_path, 
    doc_type="BB", 
    meeting_title="",
    doc_number="2710",
    signer_title="",
    signer_name="",
    recipients=None
):
    doc = docx.Document()

    # 1. Page margins (Letter 8.5 x 11.0 in, Top/Bottom 0.75", Left 1.0", Right 0.75")
    for s in doc.sections:
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.75)

    # 2. Header Table (1 dòng, 2 cột không đường viền)
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    c_left = tbl.cell(0, 0)
    c_right = tbl.cell(0, 1)
    c_left.width = Inches(3.25)
    c_right.width = Inches(3.50)
    set_cell_borderless(c_left)
    set_cell_borderless(c_right)

    # Left: Đơn vị ban hành
    p = c_left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("NGÂN HÀNG TMCP ĐẦU TƯ\nVÀ PHÁT TRIỂN VIỆT NAM")
    set_run_font(r, size_pt=12.0, bold=True)

    p_l2 = c_left.add_paragraph()
    p_l2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_l2.paragraph_format.space_after = Pt(2)
    r = p_l2.add_run("————————")
    set_run_font(r, size_pt=11.0)

    so_hieu = "/BB-BIDV" if doc_type == "BB" else "/TB-BIDV"
    num = doc_number.strip() if doc_number else "2710"
    p_so = c_left.add_paragraph()
    p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_so.paragraph_format.space_after = Pt(0)
    r = p_so.add_run(f"Số:   {num}   {so_hieu}")
    set_run_font(r, size_pt=12.0)

    # Right: Quốc hiệu Tiêu ngữ & Ngày tháng
    p_r1 = c_right.paragraphs[0]
    p_r1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r1.paragraph_format.line_spacing = 1.15
    p_r1.paragraph_format.space_after = Pt(2)
    r = p_r1.add_run("CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    set_run_font(r, size_pt=12.0, bold=True)

    p_r2 = c_right.add_paragraph()
    p_r2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r2.paragraph_format.space_after = Pt(2)
    r = p_r2.add_run("Độc lập - Tự do - Hạnh phúc")
    set_run_font(r, size_pt=12.5, bold=True)

    p_r3 = c_right.add_paragraph()
    p_r3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_r3.paragraph_format.space_after = Pt(4)
    r = p_r3.add_run("————————————")
    set_run_font(r, size_pt=11.0)

    import datetime
    today = datetime.date.today()
    date_str = f"Hà Nội, ngày {today.day:02d} tháng {today.month:02d} năm {today.year}"
    p_date = c_right.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_date.paragraph_format.space_after = Pt(0)
    r = p_date.add_run(date_str)
    set_run_font(r, size_pt=12.0, italic=True)

    # Spacing after header
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(8)
    p_sp.paragraph_format.space_after = Pt(0)

    # 3. Title Block
    title_text = "BIÊN BẢN CUỘC HỌP" if doc_type == "BB" else "THÔNG BÁO KẾT LUẬN"
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t.paragraph_format.space_before = Pt(16)
    p_t.paragraph_format.space_after = Pt(6)
    r = p_t.add_run(title_text)
    set_run_font(r, size_pt=15.0, bold=True)

    default_title = "Phiên họp Cấp cao Song phương giữa Ban Lãnh đạo Ngân hàng TMCP Đầu tư và\nPhát triển Việt Nam (BIDV) và Ban Lãnh đạo Tập đoàn Tencent (Trung Quốc)\nvề Hợp tác Chiến lược Toàn diện trong lĩnh vực Chuyển đổi số, AI, Cloud và Thanh toán"
    use_title = meeting_title.strip() if meeting_title else default_title

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(4)
    r = p_sub.add_run(use_title)
    set_run_font(r, size_pt=13.5, bold=True)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_line.paragraph_format.space_after = Pt(14)
    r = p_line.add_run("————————————")
    set_run_font(r, size_pt=12.0)

    # 4. Parse content line by line
    raw_lines = ai_text.splitlines()
    for raw_line in raw_lines:
        line = raw_line.strip()
        if is_ignorable_line(line):
            continue

        if "Nơi nhận:" in line or "ĐẠI DIỆN ĐỒNG CHỦ TRÌ" in line or "CHỮ KÝ" in line:
            break

        # Heading mục lớn
        if line.startswith("## ") or (len(line) > 2 and line[0] in "12345" and line[1] == "." and " " in line[:4] and not line.startswith("2.1") and not line.startswith("2.2")):
            clean = line.lstrip("#").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            add_runs_to_paragraph(p, clean, default_size=13.0, force_bold=True)

        # Heading mục con
        elif line.startswith("### ") or line.startswith("2.1") or line.startswith("2.2"):
            clean = line.lstrip("#").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            add_runs_to_paragraph(p, clean, default_size=13.0, force_bold=True)

        # 3 gạch đầu dòng chuẩn
        elif any(line.startswith(prefix) for prefix in [
            "- **- Nội dung", "- **- Đơn vị", "- **- Thời hạn",
            "- Nội dung chỉ đạo", "- Đơn vị thực hiện", "- Thời hạn hoàn thành",
            "**- Nội dung", "**- Đơn vị", "**- Thời hạn",
            "* **- Nội dung", "* **- Đơn vị", "* **- Thời hạn",
            "* - Nội dung", "* - Đơn vị", "* - Thời hạn"
        ]):
            clean = re.sub(r"^[* \-]+", "", line).strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(0.0)
            p.paragraph_format.space_after = Pt(4)

            if ":" in clean:
                label_part, val_part = clean.split(":", 1)
                label_clean = "- " + label_part.replace("**", "").replace("*", "").strip() + ": "
                
                r_lbl = p.add_run(label_clean)
                set_run_font(r_lbl, size_pt=13.0, bold=True)
                
                is_time = "Thời hạn" in label_clean
                if is_time:
                    val_clean = val_part.replace("**", "").replace("*", "").strip()
                    r_val = p.add_run(val_clean)
                    set_run_font(r_val, size_pt=13.0, italic=True)
                else:
                    add_runs_to_paragraph(p, val_part, default_size=13.0)
            else:
                add_runs_to_paragraph(p, "- " + clean, default_size=13.0)

        # Đầu mục 1. Về ...:, 2. Về ...:
        elif re.match(r"^(\*\*)?[0-9]+\.\s+Về", line):
            clean = line.replace("**", "").strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.keep_with_next = True
            add_runs_to_paragraph(p, clean, default_size=13.0, force_bold=True)

        # 5 Luận điểm 1. Động từ: ..., 2. Động từ: ...
        elif re.match(r"^[0-9]+\.\s+", line):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)

            if ":" in line:
                pre, post = line.split(":", 1)
                r_pre = p.add_run(pre.replace("**", "").replace("*", "").strip() + ": ")
                set_run_font(r_pre, size_pt=13.0, bold=True)
                add_runs_to_paragraph(p, post, default_size=13.0)
            else:
                add_runs_to_paragraph(p, line, default_size=13.0)

        # Danh sách gạch đầu dòng thông thường
        elif line.startswith("* ") or line.startswith("- "):
            clean = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(0.0)
            p.paragraph_format.space_after = Pt(3)
            
            r_bullet = p.add_run("- ")
            set_run_font(r_bullet, size_pt=13.0)
            add_runs_to_paragraph(p, clean, default_size=13.0)

        # Đoạn văn xuôi
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.first_line_indent = Inches(0.5)
            p.paragraph_format.space_after = Pt(6)
            add_runs_to_paragraph(p, line, default_size=13.0)

    # 5. Footer Table (Nơi nhận & Chữ ký cantSplit)
    f_tbl = doc.add_table(rows=1, cols=2)
    f_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    f_tbl.autofit = False

    trPr = f_tbl.rows[0]._tr.get_or_add_trPr()
    trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))

    f_left = f_tbl.cell(0, 0)
    f_right = f_tbl.cell(0, 1)
    f_left.width = Inches(3.50)
    f_right.width = Inches(3.25)
    set_cell_borderless(f_left)
    set_cell_borderless(f_right)

    # Nơi nhận
    p_nn = f_left.paragraphs[0]
    p_nn.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_nn.paragraph_format.space_after = Pt(2)
    r = p_nn.add_run("Nơi nhận:")
    set_run_font(r, size_pt=12.0, bold=True, italic=True)

    items = recipients if (recipients and isinstance(recipients, list)) else [
        "- Lãnh đạo tham dự cuộc họp;",
        "- Ban Lãnh đạo đối tác / đơn vị liên quan;",
        "- Các Khối, Ban, Trung tâm thuộc BIDV;",
        "- Lưu: VT, VP, Tổ Thư ký đoàn."
    ]
    for it in items:
        p_it = f_left.add_paragraph()
        p_it.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_it.paragraph_format.space_after = Pt(1)
        r = p_it.add_run(it)
        set_run_font(r, size_pt=11.0)

    # Chữ ký
    p_sg = f_right.paragraphs[0]
    p_sg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sg.paragraph_format.line_spacing = 1.15
    p_sg.paragraph_format.space_after = Pt(2)

    use_signer_title = signer_title.strip() if signer_title else ("TỔNG GIÁM ĐỐC" if doc_type == "TBKL" else "ĐẠI DIỆN ĐỒNG CHỦ TRÌ PHIÊN HỌP\nTỔNG GIÁM ĐỐC BIDV")
    r = p_sg.add_run(use_signer_title)
    set_run_font(r, size_pt=12.0, bold=True)

    for _ in range(4):
        p_b = f_right.add_paragraph()
        p_b.paragraph_format.space_after = Pt(0)
        p_b.paragraph_format.line_spacing = 1.0

    use_signer_name = signer_name.strip() if signer_name else "HOÀNG VIỆT HÙNG"
    p_nm = f_right.add_paragraph()
    p_nm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_nm.paragraph_format.space_after = Pt(0)
    r = p_nm.add_run(use_signer_name)
    set_run_font(r, size_pt=12.0, bold=True)

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)

    # Convert to PDF
    pdf_path = os.path.splitext(docx_path)[0] + ".pdf"
    try:
        subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", docx_path, "--outdir", os.path.dirname(docx_path)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except Exception:
        pass

    return docx_path, pdf_path if os.path.exists(pdf_path) else None
