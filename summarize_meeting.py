import json
import os
import sys
import time
import urllib.request
import docx

INPUT_FILE = "/home/hieu/Downloads/2026_10_27 Ban_ghi_chep_Phien_hop_cap_cao_BIDV_Tencent_20261027_Phan_vai_chuan.docx"
OUTPUT_MD = "/home/hieu/Downloads/Bien_ban_hop_song_phuong_BIDV_Tencent_20261027.md"
OUTPUT_DOCX = "/home/hieu/Downloads/Bien_ban_hop_song_phuong_BIDV_Tencent_20261027.docx"

print(f"Reading input docx: {INPUT_FILE}...")
doc = docx.Document(INPUT_FILE)
lines = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)

for tbl in doc.tables:
    for row in tbl.rows:
        row_str = " | ".join(c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip())
        if row_str:
            lines.append(row_str)

full_content = "\n".join(lines)
print(f"Loaded {len(lines)} lines, {len(full_content)} chars.")

system_prompt = """Bạn là Thư ký cấp cao của Ngân hàng TMCP Đầu tư và Phát triển Việt Nam (BIDV). Nhiệm vụ của bạn là đọc toàn bộ biên bản ghi chép chi tiết, phân vai của Phiên họp cấp cao song phương giữa Ban Lãnh đạo BIDV và Tập đoàn Tencent (Trung Quốc) ngày 27/10/2026 tại Trụ sở Tencent (Thâm Quyến) và soạn thảo thành một **BIÊN BẢN HỌP SONG PHƯƠNG** chính thức, chuẩn mực, chuyên nghiệp và có giá trị pháp lý/vận hành cao.

Yêu cầu về văn phong và cấu trúc:
- Văn phong ngoại giao, tài chính - ngân hàng cao cấp, trang trọng, cô đọng, rõ ràng.
- Đầy đủ các phần:
  I. THÔNG TIN CHUNG (Thời gian, địa điểm, thành phần đoàn BIDV do TGĐ Hoàng Việt Hùng dẫn đầu cùng 5 PTGĐ và lãnh đạo các Ban; thành phần Tencent; Chủ trì & Thư ký).
  II. MỤC ĐÍCH & Ý NGHĨA CUỘC HỌP.
  III. NỘI DUNG TRAO ĐỔI VÀ THỐNG NHẤT CỦA HAI BÊN:
      1. Tầm nhìn chiến lược và chuyển đổi số toàn diện.
      2. Ứng dụng Trí tuệ nhân tạo (AI): Chiến lược chuyển đổi AI, AI Platform, AI Governance, Multi-Agent, tuân thủ Luật AI tại Việt Nam, các bài toán Quick-win.
      3. Hạ tầng Đám mây (Cloud), Vận hành tự động (AIOps) và An ninh bảo mật.
      4. Định danh số (eKYC), Sinh trắc học và Thanh toán thế hệ mới (Pay by Palm).
      5. Định hướng hợp tác tài chính, đầu tư và cơ chế phối hợp dài hạn.
      6. Lễ Ký kết Biên bản ghi nhớ (MOU) và Trao quà lưu niệm.
  IV. KẾT LUẬN & KẾ HOẠCH HÀNH ĐỘNG (ACTION ITEMS) - có bảng phân công công việc, đơn vị đầu mối hai bên và mốc thời gian thực hiện.
  V. ĐIỀU KHOẢN HIỆU LỰC & CHỮ KÝ XÁC NHẬN CỦA ĐẠI DIỆN HAI BÊN."""

user_content = f"""Dưới đây là toàn bộ bản ghi chép chi tiết phiên họp:

---
{full_content}
---

Hãy soạn thảo BIÊN BẢN HỌP SONG PHƯƠNG chính thức theo đúng các yêu cầu trên."""

req_data = {
    "model": "gemini-argon",
    "stream": True,
    "messages": [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_content}
    ]
}

print("Sending request to gemini-stealth-bridge (gemini-argon)...")
req = urllib.request.Request(
    "http://127.0.0.1:20133/v1/chat/completions",
    data=json.dumps(req_data).encode("utf-8"),
    headers={
        "Authorization": "Bearer public",
        "Content-Type": "application/json"
    }
)

start_time = time.time()
full_response = []
chars_count = 0

with urllib.request.urlopen(req, timeout=300) as resp:
    for line in resp:
        line = line.decode("utf-8").strip()
        if not line or line.startswith(": ping"):
            continue
        if line.startswith("data: "):
            payload = line[6:]
            if payload == "[DONE]":
                break
            try:
                data = json.loads(payload)
                delta = data["choices"][0]["delta"]
                text = delta.get("content", "")
                if text:
                    full_response.append(text)
                    chars_count += len(text)
                    sys.stdout.write(f"\rReceived {chars_count} characters... ({time.time() - start_time:.1f}s)")
                    sys.stdout.flush()
            except Exception:
                pass

print(f"\nCompleted in {time.time() - start_time:.1f}s!")
full_text = "".join(full_response)

with open(OUTPUT_MD, "w", encoding="utf-8") as f:
    f.write(full_text)
print(f"Saved Markdown to {OUTPUT_MD}")

# Now build the DOCX file
print(f"Generating DOCX: {OUTPUT_DOCX}...")
out_doc = docx.Document()

# Page margins: standard 1 inch
for section in out_doc.sections:
    section.top_margin = docx.shared.Inches(1)
    section.bottom_margin = docx.shared.Inches(1)
    section.left_margin = docx.shared.Inches(1)
    section.right_margin = docx.shared.Inches(1)

for line in full_text.splitlines():
    stripped = line.strip()
    if not stripped:
        continue
    if stripped.startswith("# "):
        h = out_doc.add_heading(stripped[2:], level=1)
    elif stripped.startswith("## "):
        h = out_doc.add_heading(stripped[3:], level=2)
    elif stripped.startswith("### "):
        h = out_doc.add_heading(stripped[4:], level=3)
    elif stripped.startswith("- ") or stripped.startswith("* "):
        p = out_doc.add_paragraph(stripped[2:], style='List Bullet')
    else:
        p = out_doc.add_paragraph(stripped)

out_doc.save(OUTPUT_DOCX)
print(f"Saved DOCX to {OUTPUT_DOCX}")
print("ALL DONE!")
